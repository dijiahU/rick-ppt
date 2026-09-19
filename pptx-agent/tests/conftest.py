from pathlib import Path
import sys

sys.path.append(str(Path(__file__).resolve().parents[1] / "skills/pptx/scripts"))

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

from pptx_core.package import unpack


@pytest.fixture
def deck(tmp_path):
    image = tmp_path / "picture.png"
    Image.new("RGB", (120, 80), "#168aad").save(image)
    pres = Presentation()
    for n in range(3):
        slide = pres.slides.add_slide(pres.slide_layouts[5])
        slide.shapes.title.text = f"Original title {n + 1}"
        for i in range(3):
            box = slide.shapes.add_textbox(Inches(0.8 + i * 2.8), Inches(2), Inches(2), Inches(1))
            box.text = f"Node {i + 1}"
        slide.shapes.add_picture(str(image), Inches(1), Inches(4), Inches(1.5))
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("Values", [3, 7])
    pres.slides[1].shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(4), Inches(3.5), Inches(4), Inches(3), data)
    table = pres.slides[2].shapes.add_table(2, 2, Inches(4), Inches(4), Inches(3), Inches(1)).table
    table.cell(0, 0).text = "Native table"
    group = pres.slides[2].shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(3), Inches(2), Inches(0.5)).text = "Grouped text"
    dest = tmp_path / "research deck.pptx"
    pres.save(dest)
    return dest


@pytest.fixture
def ws(deck):
    return unpack(deck)
