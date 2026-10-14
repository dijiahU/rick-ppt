import hashlib
import io
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from types import SimpleNamespace
from unittest.mock import Mock, patch
import zipfile

from lxml import etree as E
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'pptx-agent/skills/pptx/scripts'))
from native_builds import set_builds
from pptx_core.package import unpack, pack, Workspace
from recovery_workspace import relocate_task
from repair_bootstrap import bootstrap_repair_workspace, RepairBootstrapError

NS = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
PYTHON = '/Users/rick/Desktop/ppt/pptx-agent/.venv/bin/python'


def digest(body):
    return hashlib.sha256(body).hexdigest()


class RepairBootstrapTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.base = Path(self.temp.name).resolve()
        self.job = self.base / 'task'; self.job.mkdir()
        self.cfg = {'python': PYTHON, 'plugin': str(ROOT / 'pptx-agent')}
        # Actual editable text, chart/workbook relationships, theme/font parts,
        # speaker notes, and native timing are all present in the frozen fixture.
        prs = Presentation(ROOT / 'pptx-agent/skills/pptx/assets/blank.pptx')
        slide = prs.slides[0]
        title = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        title.text = 'Frozen current label'
        title.text_frame.paragraphs[0].runs[0].font.name = 'Noto Sans CJK SC'
        data = ChartData(); data.categories = ['A', 'B']; data.add_series('Evidence', [2, 3])
        chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
                                      Inches(1), Inches(2), Inches(5), Inches(3), data)
        slide.notes_slide.notes_text_frame.text = 'Persistent context; then reveal chart.'
        source = self.base / 'current.pptx'; prs.save(source)
        ws = unpack(source, self.base / 'fixture')
        part = ws.root / 'ppt/slides/slide1.xml'
        tree = E.parse(str(part)); set_builds(tree.getroot(), [[title.shape_id], [chart.shape_id]])
        part.write_bytes(E.tostring(tree, xml_declaration=True, encoding='UTF-8', standalone=True))
        self.frozen = self.base / 'frozen.pptx'; pack(ws.root, self.frozen)
        self.artifact = self.frozen.read_bytes()
        self.identity = {'artifact': self.artifact, 'artifact_sha256': digest(self.artifact),
                         'task_id': 'fixture-task', 'input_revision': 3,
                         'plugin_version': 'fixture-runtime', 'repair_round': 1}
        self.calls = 0

    def tearDown(self):
        self.temp.cleanup()

    def unpack_callback(self, source, base):
        self.calls += 1
        ws = unpack(source, base)
        return {'workspace': str(ws.root), 'state': str(ws.state_path)}

    def bootstrap(self, **extra):
        values = dict(self.identity); values.update(extra)
        return bootstrap_repair_workspace(self.job, self.cfg, unpack=self.unpack_callback, **values)

    def files(self, path):
        return {p.relative_to(path).as_posix(): digest(p.read_bytes())
                for p in path.rglob('*') if p.is_file()}

    def test_default_canonical_cli_preserves_every_native_part_and_history(self):
        historical = self.job / 'old-static.pptx'
        historical.write_bytes((ROOT / 'pptx-agent/skills/pptx/assets/blank.pptx').read_bytes())
        before = historical.read_bytes()
        result = bootstrap_repair_workspace(self.job, self.cfg, **self.identity)
        workspace = Path(result['workspace'])
        with zipfile.ZipFile(self.frozen) as archive:
            for name in archive.namelist():
                if not name.endswith('/'):
                    self.assertEqual((workspace / name).read_bytes(), archive.read(name), name)
        tree = E.parse(str(workspace / 'ppt/slides/slide1.xml'))
        self.assertEqual(len(tree.xpath('.//p:cTn[@presetClass="entr"]', namespaces=NS)), 2)
        self.assertIn('Noto Sans CJK SC', (workspace / 'ppt/slides/slide1.xml').read_text())
        self.assertTrue(list((workspace / 'ppt/embeddings').glob('*.xlsx')))
        self.assertTrue((workspace / 'ppt/slides/_rels/slide1.xml.rels').is_file())
        self.assertEqual(Path(result['candidate_path']).read_bytes(), self.artifact)
        self.assertEqual(historical.read_bytes(), before)
        self.assertEqual(self.frozen.read_bytes(), self.artifact)
        Workspace(workspace.parent).check_original()

    def test_same_binding_retains_legitimate_partial_edits_without_reunpack(self):
        first = self.bootstrap(); workspace = Path(first['workspace'])
        slide = workspace / 'ppt/slides/slide1.xml'
        changed = slide.read_bytes().replace(b'Frozen current label', b'Corrected partial label')
        slide.write_bytes(changed)
        # Current changes need not match the frozen baseline, and refresh may
        # alter the active state. Bootstrap must not reset either to the seed.
        ws = Workspace(workspace.parent); ws.refresh()
        before = self.files(workspace.parent)
        again = self.bootstrap()
        self.assertTrue(again['reused']); self.assertTrue(again['partial_edits_retained'])
        self.assertEqual(again['workspace'], first['workspace']); self.assertEqual(self.calls, 1)
        self.assertEqual(self.files(workspace.parent), before)
        self.assertEqual(slide.read_bytes(), changed)

    def test_expected_hash_mismatch_rejects_before_creating_workspace(self):
        before = self.files(self.job)
        with self.assertRaisesRegex(RepairBootstrapError, 'hash mismatch'):
            self.bootstrap(artifact_sha256='0' * 64)
        self.assertEqual(self.files(self.job), before)

    def test_tampered_original_or_baseline_is_not_discarded_or_silently_reseeded(self):
        first = self.bootstrap(); home = Path(first['workspace']).parent
        state_path = home / 'state.json'; state = json.loads(state_path.read_text())
        state['baseline']['ppt/slides/slide1.xml']['hash'] = '0' * 64
        state_path.write_text(json.dumps(state))
        before = self.files(Path(first['home']))
        with self.assertRaisesRegex(RepairBootstrapError, 'baseline'):
            self.bootstrap()
        self.assertEqual(self.files(Path(first['home'])), before); self.assertEqual(self.calls, 1)
        state['baseline']['ppt/slides/slide1.xml']['hash'] = first['baseline_parts']['ppt/slides/slide1.xml']['hash']
        state_path.write_text(json.dumps(state))
        original = home / 'original.pptx'; original.chmod(0o644); original.write_bytes(self.artifact + b'tamper')
        with self.assertRaisesRegex(RepairBootstrapError, 'candidate bytes changed'):
            self.bootstrap()
        self.assertEqual(original.read_bytes(), self.artifact + b'tamper')

    def test_owner_revision_runtime_round_bindings_keep_distinct_history(self):
        first = self.bootstrap(); before = self.files(Path(first['home']))
        for field, value in [('task_id', 'different-owner'), ('input_revision', 4),
                             ('plugin_version', 'next-runtime'), ('repair_round', 2)]:
            other = self.bootstrap(**{field: value})
            self.assertNotEqual(other['generation'], first['generation'])
            self.assertFalse(other['reused'])
            self.assertEqual(self.files(Path(first['home'])), before)
        receipt = Path(first['receipt_path']); value = json.loads(receipt.read_text())
        value['binding']['task_id'] = 'forged-owner'; receipt.write_text(json.dumps(value))
        with self.assertRaisesRegex(RepairBootstrapError, 'binding disagrees'):
            self.bootstrap()

    def test_recovery_relocation_retains_binding_and_partial_source(self):
        first = self.bootstrap(); workspace = Path(first['workspace'])
        slide = workspace / 'ppt/slides/slide1.xml'
        slide.write_bytes(slide.read_bytes().replace(b'Frozen current label', b'Restored partial label'))
        recovered = self.base / 'restored-task'; shutil.copytree(self.job, recovered)
        relocate_task(self.job, recovered)
        result = bootstrap_repair_workspace(recovered, self.cfg, unpack=self.unpack_callback, **self.identity)
        self.assertTrue(result['reused']); self.assertEqual(self.calls, 1)
        self.assertIn(b'Restored partial label', (Path(result['workspace']) / 'ppt/slides/slide1.xml').read_bytes())
        self.assertEqual(Path(result['candidate_path']).read_bytes(), self.artifact)

    def test_incomplete_unpack_and_symlinks_preserve_evidence_and_refuse_reuse(self):
        def fail(source, base):
            raise RuntimeError('simulated interrupted unpack')
        with self.assertRaisesRegex(RuntimeError, 'interrupted unpack'):
            bootstrap_repair_workspace(self.job, self.cfg, unpack=fail, **self.identity)
        before = self.files(self.job)
        with self.assertRaisesRegex(RepairBootstrapError, 'Incomplete'):
            self.bootstrap()
        self.assertEqual(self.files(self.job), before)
        clean = self.base / 'clean-task'; clean.mkdir()
        (clean / 'repair-workspaces').symlink_to(self.base)
        with self.assertRaisesRegex(RepairBootstrapError, 'symlinks'):
            bootstrap_repair_workspace(clean, self.cfg, unpack=self.unpack_callback, **self.identity)

    def test_out_of_scope_or_mutating_unpack_callback_is_not_accepted(self):
        def outside(source, base):
            return {'workspace': str(self.base / 'outside/workspace')}
        with self.assertRaisesRegex(RepairBootstrapError, 'out-of-scope'):
            bootstrap_repair_workspace(self.job, self.cfg, unpack=outside, **self.identity)
        clean = self.base / 'mutating-task'; clean.mkdir()
        def mutate(source, base):
            result = self.unpack_callback(source, base)
            part = Path(result['workspace']) / 'ppt/slides/slide1.xml'
            part.write_bytes(part.read_bytes().replace(b'Frozen current label', b'Unexpected rewrite'))
            return result
        with self.assertRaisesRegex(RepairBootstrapError, 'Unpack changed'):
            bootstrap_repair_workspace(clean, self.cfg, unpack=mutate, **self.identity)

    def test_unsafe_zip_members_refuse_before_any_bootstrap_write(self):
        for name in ('../outside.xml', '/absolute.xml', 'ppt\\escape.xml'):
            body = io.BytesIO()
            with zipfile.ZipFile(body, 'w') as archive:
                archive.writestr('[Content_Types].xml', b'<Types/>')
                archive.writestr('ppt/presentation.xml', b'<presentation/>')
                archive.writestr(name, b'unsafe')
            candidate = body.getvalue(); before = self.files(self.job)
            with self.assertRaisesRegex(RepairBootstrapError, 'Unsafe'):
                self.bootstrap(artifact=candidate, artifact_sha256=digest(candidate))
            self.assertEqual(self.files(self.job), before)

    def test_active_workspace_symlink_is_refused_without_erasing_partial_source(self):
        result = self.bootstrap(); workspace = Path(result['workspace'])
        private = self.base / 'outside.xml'; private.write_bytes(b'outside must remain')
        (workspace / 'unsafe.xml').symlink_to(private)
        with self.assertRaisesRegex(RepairBootstrapError, 'symlinks'):
            self.bootstrap()
        self.assertTrue((workspace / 'unsafe.xml').is_symlink())
        self.assertEqual(private.read_bytes(), b'outside must remain')

    def host_execution(self):
        from workflow import Execution
        host = self.base / 'host'; host.mkdir(exist_ok=True)
        bridge = SimpleNamespace(ROOT=host, request=lambda *a, **kw: {'ok': True},
                                 sandbox=lambda cfg, root, args: args)
        cfg = {**self.cfg, 'plugin_version': self.identity['plugin_version']}
        task = {'id': self.identity['task_id'], 'input_revision': self.identity['input_revision']}
        run = Execution(bridge, cfg, task, self.job, SimpleNamespace(check=lambda: None),
                        SimpleNamespace(), trace=Mock())

        def run_check(root, command, **kwargs):
            self.assertEqual(root, self.job)
            self.assertEqual(command[0], PYTHON)
            self.assertIn('unpack', command)
            return subprocess.run(command, cwd=root, capture_output=True, text=True,
                                  timeout=30, check=True).stdout
        run.run_check = Mock(side_effect=run_check)
        return run

    def test_host_seeds_matching_frozen_candidate_instead_of_old_delivery_workspace(self):
        run = self.host_execution()
        static = self.job / 'old-static.pptx'
        static.write_bytes((ROOT / 'pptx-agent/skills/pptx/assets/blank.pptx').read_bytes())
        old = unpack(static, self.job / 'old-author')
        (self.job / 'delivery.json').write_text(json.dumps({'path': str(static), 'workspace': str(old.root)}))
        old_before = self.files(old.home); static_before = static.read_bytes()
        findings = self.job / 'findings.json'
        findings.write_text(json.dumps({'artifact_sha256': self.identity['artifact_sha256']}))
        run.frozen_candidate = {'artifact': self.artifact,
                                'sha256': self.identity['artifact_sha256'],
                                'input_revision': self.identity['input_revision'],
                                'plugin_version': self.identity['plugin_version']}
        # Exercise Execution's real callback and real bootstrap. Tool manifest
        # preparation belongs to another unit and must not mask seed assertions.
        with patch('tool_preparation.prepare_tool_manifest') as prepare:
            result = run.seed_repair_workspace(findings, 1)
        self.assertEqual(result['binding']['artifact_sha256'], self.identity['artifact_sha256'])
        self.assertEqual(Path(result['candidate_path']).read_bytes(), self.artifact)
        self.assertNotEqual(Path(result['workspace']), old.root)
        native = E.parse(str(Path(result['workspace']) / 'ppt/slides/slide1.xml'))
        self.assertEqual(len(native.xpath('.//p:cTn[@presetClass="entr"]', namespaces=NS)), 2)
        self.assertEqual(self.files(old.home), old_before); self.assertEqual(static.read_bytes(), static_before)
        self.assertEqual(run.run_check.call_count, 1)
        prepare.assert_called_once_with(run.cfg, self.job, workspace=result['workspace'])
        self.assertEqual(json.loads((self.job / 'repair-target.json').read_text()), result)

    def test_host_rejects_review_digest_and_revision_mismatches_before_unpack(self):
        from review_sessions import ReviewSessionError
        run = self.host_execution()
        findings = self.job / 'findings.json'
        run.frozen_candidate = {'artifact': self.artifact,
                                'sha256': self.identity['artifact_sha256'],
                                'input_revision': self.identity['input_revision'],
                                'plugin_version': self.identity['plugin_version']}
        findings.write_text(json.dumps({'artifact_sha256': '0' * 64}))
        with self.assertRaisesRegex(ReviewSessionError, 'verified frozen candidate'):
            run.seed_repair_workspace(findings, 1)
        findings.write_text(json.dumps({'artifact_sha256': self.identity['artifact_sha256']}))
        run.frozen_candidate['input_revision'] += 1
        with self.assertRaisesRegex(ReviewSessionError, 'verified frozen candidate'):
            run.seed_repair_workspace(findings, 1)
        run.run_check.assert_not_called()
        self.assertFalse((self.job / 'repair-workspaces').exists())


if __name__ == '__main__':
    unittest.main()
