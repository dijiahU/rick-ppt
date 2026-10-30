"""Exact original artifacts, real native fingerprints and bounded render routing."""
import hashlib
import json
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest

from PIL import Image
from pptx import Presentation
from review_comparison import prepare, prepare_refinement

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'pptx-agent/skills/pptx/scripts'))
from pptx_core.package import unpack


class Tests(unittest.TestCase):
    def setUp(self):
        temp=tempfile.TemporaryDirectory();self.addCleanup(temp.cleanup)
        self.job=Path(temp.name).resolve();self.packet=self.job/'packet';self.packet.mkdir()
        deck=Presentation()
        for title in ('Original one','Original two'):
            slide=deck.slides.add_slide(deck.slide_layouts[5]);slide.shapes.title.text=title
        self.source=self.job/'original.pptx';deck.save(self.source)
        deck.slides[1].shapes.title.text='Revised two';candidate=self.job/'candidate.pptx';deck.save(candidate)
        self.artifact=candidate.read_bytes();self.original=self.source.read_bytes();self.calls=[]
        self.payload={'mode':'edit','attachments':[{'path':'original.pptx',
            'size':len(self.original),'sha256':hashlib.sha256(self.original).hexdigest()}]}
        def check(root,command,**kwargs):
            self.calls.append(command)
            if 'unpack' in command:
                ws=unpack(root/'original.pptx');return json.dumps({'workspace':str(ws.root)})
            page=int(command[-1]);image=root/f'rendered-{page}.png'
            Image.new('RGB',(160,90),'gold').save(image)
            return json.dumps({'pages':[str(image)]})
        self.run=SimpleNamespace(job=self.job,journal=None,cfg={'plugin':str(ROOT/'pptx-agent'),'python':sys.executable},
            lease=SimpleNamespace(check=lambda:None),run_check=check,
            bridge=SimpleNamespace(prepare=lambda cfg:Path(tempfile.mkdtemp(dir=self.job)).resolve(),
                                   sandbox=lambda cfg,root,command:command))

    def index(self):
        return json.loads((self.packet/'edit-comparison/index.json').read_text())

    def test_only_changed_dependency_page_is_rendered_and_original_is_preserved(self):
        prepare(self.run,self.artifact,self.packet,self.payload)
        data=self.index();self.assertEqual(data['status'],'prepared')
        self.assertEqual([p['dependencies_unchanged'] for p in data['pages']],[True,False])
        self.assertNotIn('before',data['pages'][0]);before=data['pages'][1]['before']
        self.assertEqual([c[-1] for c in self.calls if 'render' in c],['2'])
        image=self.packet/before['path'];self.assertEqual(hashlib.sha256(image.read_bytes()).hexdigest(),before['sha256'])
        with Image.open(image) as im:self.assertEqual(im.size,(900,506))
        self.assertEqual(self.source.read_bytes(),self.original)
        self.assertEqual(data['candidate_sha256'],hashlib.sha256(self.artifact).hexdigest())

    def test_unchanged_deck_and_ambiguous_originals_do_not_render(self):
        prepare(self.run,self.original,self.packet,self.payload)
        self.assertEqual(self.index()['status'],'prepared');self.assertEqual(self.calls,[])
        self.payload['attachments']*=2
        prepare(self.run,self.artifact,self.packet,self.payload)
        self.assertEqual(self.index()['status'],'no_unique_original_pptx');self.assertEqual(self.calls,[])

    def test_source_checksum_mismatch_cannot_become_comparison_evidence(self):
        self.payload['attachments'][0]['sha256']='0'*64
        with self.assertRaisesRegex(ValueError,'checksum changed'):
            prepare(self.run,self.artifact,self.packet,self.payload)
        self.assertEqual(self.calls,[])

    def test_renderer_failure_keeps_explicit_unavailable_state(self):
        self.run.run_check=lambda *a,**k:(_ for _ in ()).throw(TimeoutError('fixture'))
        prepare(self.run,self.artifact,self.packet,self.payload)
        self.assertTrue(self.index()['status'].startswith('preview_unavailable'))
        self.assertEqual(self.index()['pages'][1]['dependencies_unchanged'],False)
        self.assertNotIn('before',self.index()['pages'][1])

    def test_new_deck_does_not_prepare_source_comparison(self):
        prepare(self.run,self.artifact,self.packet,{**self.payload,'mode':'create'})
        self.assertFalse((self.packet/'edit-comparison').exists())

    def test_presentation_size_change_cannot_be_reported_as_unchanged_pages(self):
        import io
        import zipfile
        from lxml import etree
        stream=io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(self.original)) as source,zipfile.ZipFile(stream,'w') as target:
            for name in source.namelist():
                body=source.read(name)
                if name=='ppt/presentation.xml':
                    tree=etree.fromstring(body)
                    tree.find('{*}sldSz').set('cx','10000000')
                    body=etree.tostring(tree)
                target.writestr(name,body)
        prepare(self.run,stream.getvalue(),self.packet,self.payload)
        self.assertEqual([p['dependencies_unchanged'] for p in self.index()['pages']],[False,False])
        self.assertFalse(self.index()['presentation_and_font_parts_unchanged'])

    def test_stop_is_not_swallowed_as_optional_preview_failure(self):
        from resilience import TaskPaused
        self.run.run_check=lambda *a,**k:(_ for _ in ()).throw(TaskPaused('fixture'))
        with self.assertRaises(TaskPaused):prepare(self.run,self.artifact,self.packet,self.payload)


class RefinementTests(unittest.TestCase):
    setUp = Tests.setUp
    def baseline(self, completion=False):
        from design_feedback import save_baseline, POLICY
        temporary = tempfile.TemporaryDirectory(); self.addCleanup(temporary.cleanup)
        self.run.records = Path(temporary.name).resolve()
        self.run.task = {'id': 'refinement-comparison-fixture', 'input_revision': 0}
        images = []
        for page in (1, 2):
            path = self.job / f'baseline-{page}.png'
            Image.new('RGB', (160, 90), 'blue').save(path); images.append(str(path))
        receipt = {'artifact_sha256': hashlib.sha256(self.original).hexdigest(),
            'round': 1, 'input_revision': 0, 'plugin_version': 'fixture-v1',
            'review_mode': 'substantive', 'designer_response': dict(POLICY),
            'content': {'findings': []}, 'visual': {'findings': [{'id': 'V1',
                'severity': 'suggestion', 'pages': [2], 'observation': 'Main relationship is detached.',
                'impact': 'Cross-reading is needed.', 'recommendation': 'Connect the subject and explanation.',
                'route': 'retain_suggestion'}]}}
        if completion:
            from design_completion import CONTRACT
            receipt['design_completion']=dict(CONTRACT)
            receipt['visual']['design_assessment']={'readiness':'needs_revision','primary_findings':['V1']}
        return save_baseline(self.run, self.original, receipt, {'pages': images}, {})

    def test_selection_compares_changed_pages_outside_original_feedback(self):
        import io
        self.baseline(completion=True)
        deck=Presentation(io.BytesIO(self.original));deck.slides[0].shapes.title.text='Changed outside feedback'
        body=io.BytesIO();deck.save(body)
        prepare_refinement(self.run,body.getvalue(),self.packet)
        data=json.loads((self.packet/'refinement-comparison/index.json').read_text())
        self.assertEqual([p['page'] for p in data['pages']],[1,2])
        self.assertEqual([p['native_dependencies_unchanged'] for p in data['pages']],[False,True])
        self.assertEqual(self.calls,[])  # Reuses verified baseline PNGs, without rerendering originals.

    def test_shared_dimension_changes_include_all_previous_pages(self):
        import io
        self.baseline(completion=True)
        deck=Presentation(io.BytesIO(self.original));deck.slide_width+=100000
        body=io.BytesIO();deck.save(body)
        prepare_refinement(self.run,body.getvalue(),self.packet)
        data=json.loads((self.packet/'refinement-comparison/index.json').read_text())
        self.assertEqual([p['page'] for p in data['pages']],[1,2])
        self.assertFalse(data['presentation_and_font_parts_unchanged'])
        self.assertTrue(all(not p['native_dependencies_unchanged'] for p in data['pages']))

    def test_refinement_reuses_only_verified_target_previews_without_rendering(self):
        baseline = self.baseline()
        prepare_refinement(self.run, self.artifact, self.packet)
        data = json.loads((self.packet / 'refinement-comparison/index.json').read_text())
        self.assertEqual(data['source_sha256'], baseline['receipt']['artifact_sha256'])
        self.assertEqual(data['candidate_sha256'], hashlib.sha256(self.artifact).hexdigest())
        self.assertEqual([p['page'] for p in data['pages']], [2])
        self.assertFalse(data['pages'][0]['native_dependencies_unchanged'])
        self.assertEqual(data['findings'], baseline['receipt']['visual']['findings'])
        info = data['pages'][0]['before']; body = (self.packet / info['path']).read_bytes()
        self.assertEqual(hashlib.sha256(body).hexdigest(), info['sha256'])
        self.assertEqual(self.calls, [])

    def test_noop_and_new_revision_cannot_publish_a_stale_baseline(self):
        self.baseline()
        prepare_refinement(self.run, self.original, self.packet)
        self.assertFalse((self.packet / 'refinement-comparison').exists())
        self.run.task['input_revision'] = 1
        prepare_refinement(self.run, self.artifact, self.packet)
        self.assertFalse((self.packet / 'refinement-comparison').exists())

    def test_changed_baseline_pixels_are_rejected(self):
        baseline = self.baseline()
        path = Path(baseline['home']) / baseline['pages'][1]['file']
        path.chmod(0o600); path.write_bytes(b'corrupt')
        with self.assertRaisesRegex(ValueError, 'bytes changed'):
            prepare_refinement(self.run, self.artifact, self.packet)


if __name__=='__main__':unittest.main()
