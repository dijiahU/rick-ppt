"""Real PPTX/pixels and mock stdio transport; no models, accounts or queue."""
import hashlib
import importlib.util
import json
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch

from PIL import Image
from pptx import Presentation
from repair_inputs import prepare, deadline_handoff
from design_feedback import save_baseline, POLICY
from durable import AppServer
from conversation import RevisionPending
from resilience import TaskPaused

ROOT = Path(__file__).resolve().parents[1]
sys.path.append(str(ROOT/'workflow'))
spec = importlib.util.spec_from_file_location('repair_protocol_fixture', ROOT/'workflow/test_app_server.py')
fixture = importlib.util.module_from_spec(spec);spec.loader.exec_module(fixture)


class Tests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory();self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name).resolve();self.job = self.root/'task';self.job.mkdir()
        self.records = self.root/'records';self.records.mkdir()
        self.cfg = {'plugin':str(ROOT/'pptx-agent'), 'python':sys.executable}
        self.run = SimpleNamespace(cfg=self.cfg, task={'id':'repair-fixture','input_revision':0},
                                   journal=None, records=self.records, job=self.job)
        deck = Presentation()
        for title in ('Question','Relation'):
            slide=deck.slides.add_slide(deck.slide_layouts[5]);slide.shapes.title.text=title
        deck.save(self.job/'baseline.pptx');self.original=(self.job/'baseline.pptx').read_bytes()
        deck.slides[1].shapes.title.text='Improved relation';deck.save(self.job/'candidate.pptx')
        self.candidate=(self.job/'candidate.pptx').read_bytes()
        self.binding={'sha256':hashlib.sha256(self.original).hexdigest(),'input_revision':0}
        (self.job/'delivery.json').write_text(json.dumps({'path':'candidate.pptx'}))
        images=[]
        for n in (1,2):
            p=self.root/f'page-{n}.png';Image.new('RGB',(1334,750),'blue').save(p);images.append(str(p))
        self.review={'artifact_sha256':self.binding['sha256'],'input_revision':0,'round':1,
            'plugin_version':'fixture-v1','review_mode':'substantive','designer_response':dict(POLICY),
            'content':{'findings':[]},'visual':{'findings':[{'id':'V1','severity':'suggestion','pages':[2],
                'observation':'Subject and explanation are detached.','impact':'Cross-reading is needed.',
                'recommendation':'Connect them while preserving qualifications and clicks.','route':'retain_suggestion'}]}}
        self.baseline=save_baseline(self.run,self.original,self.review,{'pages':images},{})

    def server(self):
        server=AppServer(cwd=self.job,config={},command=[sys.executable,'-u','-c',fixture.FAKE_SERVER])
        server.__enter__();self.addCleanup(server.__exit__,None,None,None)
        thread=server.start_thread()['thread']['id'];turn=server.start_turn(thread,'WAIT')['id']
        return server,thread,turn

    def handoff(self, server=None, check=None):
        server,thread,turn=server or self.server()
        return deadline_handoff(server,thread,turn,self.job,self.binding,
            check_current=check or (lambda revision:None),timeout=2)

    def test_declared_primary_gap_controls_direct_repair_input_not_json_order(self):
        from design_completion import CONTRACT
        self.run.task['id']='repair-primary-priority-fixture'
        self.review['design_completion']=dict(CONTRACT)
        second={**self.review['visual']['findings'][0],'id':'V2','pages':[1],
                'observation':'The principal example is too small to inspect.'}
        self.review['visual']['findings'].append(second)
        self.review['visual']['design_assessment']={'readiness':'needs_revision','primary_findings':['V2']}
        source=[str(Path(self.baseline['home'])/p['file']) for p in self.baseline['pages']]
        self.baseline=save_baseline(self.run,self.original,self.review,{'pages':source},{})
        text,images=prepare(self.run,self.review)
        payload=json.loads(text[text.index('{"artifact_sha256"'):])
        self.assertEqual(payload['primary_finding']['id'],'V2')
        self.assertEqual([p['page'] for p in payload['pages']],[1])
        self.assertEqual([p.name for p in images],['page-1.png'])
        self.assertEqual(payload['pages'][0]['reviewed_image_sha256'],self.baseline['pages'][0]['sha256'])
        self.assertEqual(payload['pages'][0]['sha256'],hashlib.sha256(images[0].read_bytes()).hexdigest())

    def test_primary_review_pixels_and_feedback_are_supplied_without_rendering(self):
        text,images=prepare(self.run,self.review)
        self.assertEqual(len(images),1);self.assertIn('V1',text)
        self.assertIn('primary_finding',text)
        self.assertTrue(images[0].is_relative_to(self.job))
        with Image.open(images[0]) as image:self.assertEqual(image.size,(900,506))
        self.assertIn(hashlib.sha256(images[0].read_bytes()).hexdigest(),text)
        self.assertNotIn(str(self.records),text)
        self.assertEqual(prepare(self.run,self.review)[1],images)

    def test_content_priority_sends_its_actual_page_and_survives_comparison(self):
        from design_completion import CONTRACT, PRESENTATION, comparison_input
        from review_comparison import prepare_refinement
        from workflow import _repair_candidate
        self.run.task['id']='content-priority-pixels'
        self.review.update(design_completion=dict(CONTRACT),content_presentation=dict(PRESENTATION))
        self.review['visual']['design_assessment']={'readiness':'ready','primary_findings':[]}
        self.review['content']={'findings':[{**self.review['visual']['findings'][0],
            'pages':[1],'observation':'The qualifier is hidden in tiny credits.'}], 'presentation_findings':['V1']}
        source=[str(Path(self.baseline['home'])/p['file']) for p in self.baseline['pages']]
        # Distinct page pixels catch accidental attachment of the visual review's page 2.
        first=self.root/'content-page.png';Image.new('RGB',(1334,750),'gold').save(first);source[0]=str(first)
        self.baseline=save_baseline(self.run,self.original,self.review,{'pages':source},{})
        self.run.author_thread='continuing-author';self.run.recovery_plan={}
        self.run.phase=Mock();self.run.complete_phase=Mock()
        self.run.delivery=lambda:json.loads((self.job/'delivery.json').read_text())
        findings=self.job/'review.json';findings.write_text(json.dumps(self.review))
        _repair_candidate(self.run,'Host scope','Continue author',findings,1,content_required=False)
        call=self.run.phase.call_args
        self.assertEqual(call.kwargs['thread'],'continuing-author')
        self.assertEqual(call.kwargs['timeout'],360)
        self.assertEqual([p.name for p in call.kwargs['images']],['page-1.png'])
        with Image.open(call.kwargs['images'][0]) as image:
            self.assertEqual(image.getpixel((0,0)),(255,215,0))
        self.assertIn('content:V1',call.args[1])
        packet=self.root/'packet';packet.mkdir()
        prepare_refinement(self.run,self.candidate,packet)
        index=json.loads((packet/'refinement-comparison/index.json').read_text())
        self.assertEqual([f['id'] for f in index['findings']],['content:V1','visual:V1'])
        self.assertEqual([p['page'] for p in index['pages']],[1,2])
        images=[]
        for row in index['pages']:
            body=(packet/row['before']['path']).read_bytes()
            self.assertEqual(hashlib.sha256(body).hexdigest(),row['before']['sha256'])
            images.extend([row['before']['path'],f"reading-page-{row['page']}.png"])
        binding=comparison_input(packet,images,hashlib.sha256(self.candidate).hexdigest())
        self.assertEqual(binding['finding_ids'],['content:V1','visual:V1'])
        self.assertTrue(binding['visual_input_complete'])

    def test_stale_review_and_no_visual_findings_do_not_invent_inputs(self):
        self.run.task['input_revision']=1
        self.assertEqual(prepare(self.run,self.review),('',[]))
        self.run.task['input_revision']=0
        self.review['visual']['findings']=[]
        self.assertEqual(prepare(self.run,self.review),('',[]))

    def test_author_symlink_cannot_redirect_host_picture_writes(self):
        outside=self.root/'outside';outside.mkdir()
        (self.job/'repair-inputs').symlink_to(outside,target_is_directory=True)
        with self.assertRaises((OSError,ValueError)):prepare(self.run,self.review)
        self.assertEqual(list(outside.iterdir()),[])

    def test_real_stdio_interruption_precedes_changed_candidate_handoff(self):
        server,thread,turn=self.server();checks=[]
        def check(revision):
            checks.append(thread in server.active_turns);self.assertEqual(revision,0)
        result=self.handoff((server,thread,turn),check)
        self.assertEqual(checks,[True,False,False])
        self.assertEqual(result['sha256'],hashlib.sha256(self.candidate).hexdigest())
        self.assertEqual(result['model_turn_status'],'interrupted')
        self.assertIn('independent reviews still required',result['acceptance'])

    def test_baseline_and_undeclared_exports_cannot_become_repair_success(self):
        (self.job/'delivery.json').write_text(json.dumps({'path':'baseline.pptx'}))
        self.assertIsNone(self.handoff())
        (self.job/'delivery.json').unlink()
        self.assertIsNone(self.handoff())

    def test_invalid_explicit_artifact_is_not_adopted(self):
        (self.job/'candidate.pptx').write_bytes(b'incomplete zip')
        self.assertIsNone(self.handoff())

    def test_stop_and_new_input_propagate_without_candidate_adoption(self):
        for exception in (TaskPaused('Owner stopped'),RevisionPending('New input')):
            with self.subTest(exception=type(exception).__name__):
                with self.assertRaises(type(exception)):
                    self.handoff(check=Mock(side_effect=exception))

    def test_candidate_changes_after_terminal_receipt_are_rejected(self):
        checks=[]
        def check(revision):
            checks.append(revision)
            if len(checks)==3:(self.job/'candidate.pptx').write_bytes(b'late write')
        with self.assertRaisesRegex(ValueError,'changed after writer interruption'):
            self.handoff(check=check)

    def test_pending_or_failed_runtime_receipt_cannot_authorize_handoff(self):
        for status in ('inProgress','failed'):
            server=SimpleNamespace(active_turns={'t':'u'},interrupt=Mock(),completed_turns={},
                wait_turn=Mock(return_value={'status':status}))
            self.assertIsNone(self.handoff((server,'t','u')))
        self.assertIsNone(self.handoff((SimpleNamespace(active_turns={},completed_turns={}), 't','u')))

    def test_native_objects_come_from_current_repair_workspace(self):
        scripts=ROOT/'pptx-agent/skills/pptx/scripts';sys.path.insert(0,str(scripts))
        from pptx_core.package import unpack
        ws=unpack(self.job/'baseline.pptx',base=self.job/'projects')
        (self.job/'repair-target.json').write_text(json.dumps({'workspace':str(ws.root)}))
        (self.job/'outline.json').write_text(json.dumps({'slides':[{'id':'one','title':'Question','summary':'Q'},
                                                                    {'id':'two','title':'Relation','summary':'R'}]}))
        (self.job/'source-notes.md').write_text('Preserve the conditions.')
        text,images=prepare(self.run,self.review)
        payload=json.loads(text[text.index('{"artifact_sha256"'):])
        self.assertEqual(payload['pages'][0]['native']['workspace'],str(ws.root))
        self.assertTrue(payload['pages'][0]['native']['objects'])
        self.assertEqual(len(images),1)

    def test_repair_call_forwards_actual_images_and_bounded_handoff_identity(self):
        from workflow import _repair_candidate
        self.run.author_thread='continuing-author';self.run.recovery_plan={}
        self.run.phase=Mock();self.run.complete_phase=Mock()
        self.run.delivery=lambda:json.loads((self.job/'delivery.json').read_text())
        findings=self.job/'review.json';findings.write_text(json.dumps(self.review))
        _repair_candidate(self.run,'Host scope','Continue the same author',findings,1,content_required=False)
        call=self.run.phase.call_args
        self.assertEqual(call.kwargs['thread'],'continuing-author')
        self.assertEqual(call.kwargs['timeout'],360)
        self.assertEqual(call.kwargs['timeout_handoff'],self.binding)
        self.assertEqual(len(call.kwargs['images']),1)
        self.assertTrue(call.kwargs['images'][0].read_bytes().startswith(b'\x89PNG\r\n\x1a\n'))
        self.assertIn('primary_finding',call.args[1])

    def test_execution_phase_accepts_deadline_candidate_only_as_handoff(self):
        from workflow import Execution
        run=Execution.__new__(Execution)
        run.cfg=self.cfg;run.job=self.job;run.task={'id':'repair-fixture','lease':'synthetic'}
        run.records=self.records;run.deadline=__import__('time').monotonic()+10
        run.journal=Mock();run.journal.state={'input_revision':0,'inbox':{},'recoveries':{}}
        run.recovery_plan={};run.conversation=Mock();run.conversation.context.return_value=None
        run.bridge=SimpleNamespace(environment=lambda root:{},render_requests=lambda *a,**k:None)
        run.reporter=Mock(last_sent=0,offset=0);run.trace=Mock();run.trajectory=None
        run.images={};run.stages=[];run.logs=[];run.threads=[];run.interactive_brokers={}
        run.media=Mock();run.fonts=Mock();run.prefetch=Mock();run.lease=Mock()
        run.last_checkpoint=__import__('time').monotonic()
        def server(**kwargs):
            kwargs['command']=[sys.executable,'-u','-c',fixture.FAKE_SERVER]
            return AppServer(**kwargs)
        with patch('workflow.AppServer',side_effect=server),patch('workflow.phase_time_limit',return_value=.15),\
             patch('workflow.AssetImporter'),patch('workflow.InteractiveHostBroker'):
            result,thread=run.phase('repair-1','WAIT',timeout=360,timeout_handoff=self.binding)
        self.assertIn('Host accepted the declared native artifact',result)
        self.assertEqual(run.stages[0]['host_handoff']['sha256'],hashlib.sha256(self.candidate).hexdigest())
        self.assertEqual(run.stages[0]['host_handoff']['model_turn_status'],'interrupted')
        run.journal.checkpoint.assert_called_once_with(self.job,reason='repair_deadline_candidate')
        run.conversation.complete_chats.assert_not_called()


if __name__=='__main__':unittest.main()
